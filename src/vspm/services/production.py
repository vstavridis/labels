"""Print & Save transaction management.

Printing and Excel are separate operations with separate statuses:

    approved row
      -> barcode reserved (registry, idempotent per row)    [label.print_status=barcode_reserved]
      -> print submitted to DYMO                            [print_unconfirmed]
      -> operator confirms the physical label               [printed_confirmed | print_failed_confirmed]
    and, after the first successful submission:
      -> Excel queued                                       [excel_status=pending|staged|queued_remote]
      -> Excel synchronized (verified)                      [synchronized] or [error|conflict]

Every step is committed to SQLite before the next one starts, so a crash or
restart resumes from the last committed step and never allocates a second
number for the same row. Reprints reuse the stored label content and barcode
and never touch Excel.
"""

from __future__ import annotations

import json
import logging
import uuid
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Optional
from zoneinfo import ZoneInfo

from vspm.barcodes.allocator import BarcodeAllocator, BarcodeError, RegistryAllocator, Reservation
from vspm.config import ExcelMode, Settings
from vspm.db.database import Database, utcnow
from vspm.db.repository import Repository
from vspm.domain.models import (
    ExcelRowPayload, ExcelStatus, LabelContent, PrintStatus, ProductionRow, RowStatus, TemplateKey,
)
from vspm.domain.validation import validate_row
from vspm.dymo.printers import LabelPrinter
from vspm.dymo.render import render_label_png
from vspm.dymo.template import TemplateRegistry
from vspm.excel.writer import ExcelSyncError, SafeWorkbookWriter
from vspm.services.transforms import TransformError, Transforms, build_excel_payload, build_label_content

log = logging.getLogger(__name__)

MESSAGES = {
    "sent": ("Η ετικέτα στάλθηκε στον εκτυπωτή. Επιβεβαιώστε ότι τυπώθηκε σωστά.",
             "Label sent to the printer. Please confirm that it printed correctly."),
    "print_failed": ("Η εκτύπωση απέτυχε: {e}", "Print failed: {e}"),
    "excel_nothing": ("Excel: δεν υπάρχει εκκρεμότητα για αυτή την ετικέτα", "Excel: nothing queued for this label"),
    "excel_already": ("Excel: έχει ήδη καταχωρηθεί", "Excel: already synchronized"),
    "excel_staged": ("Excel: η γραμμή μπήκε σε αναμονή. Καταχωρήστε την από το Ιστορικό.",
                     "Excel: row staged. Apply it from the History page (or reconcile after entering it manually)."),
    "excel_queued": ("Excel: στάλθηκε στην ουρά του υπολογιστή Excel.", "Excel: queued for the designated Excel computer."),
    "excel_written": ("Excel: γράφτηκε στη γραμμή {n} (επαληθεύτηκε).", "Excel: written to row {n} (verified)."),
    "excel_present": ("Excel: υπάρχει ήδη στη γραμμή {n} (επαληθεύτηκε).", "Excel: already present at row {n} (verified)."),
    "excel_failed": ("Το Excel ΔΕΝ ενημερώθηκε ({st}): {e}. Η ετικέτα είναι ασφαλής· δοκιμάστε ξανά αργότερα.",
                     "Excel NOT updated ({st}): {e}. The label is safe; retry later."),
    "manual_recorded": ("Η εκτύπωση επιβεβαιώθηκε από τον χειριστή.", "Print confirmed by the operator."),
    "print_uncertain": ("{b}: άγνωστο αν τυπώθηκε (η σελίδα ανανεώθηκε κατά την αποστολή). Ελέγξτε τον εκτυπωτή.",
                        "{b}: unknown whether it printed (page reloaded while sending). Check the printer."),
    "rec_queued": ("{b}: η καταχώρηση Excel μπήκε σε αναμονή μετά από επανεκκίνηση",
                   "{b}: Excel entry queued after restart"),
    "rec_reserved": ("{b}: ο αριθμός δεσμεύτηκε αλλά δεν τυπώθηκε (πατήστε ξανά Εκτύπωση & Αποθήκευση)",
                     "{b}: number reserved but not printed yet (Print & Save again to continue)"),
}


class ProductionError(Exception):
    pass


class ValidationFailed(ProductionError):
    def __init__(self, errors: dict[str, str]):
        super().__init__("; ".join(f"{k}: {v}" for k, v in errors.items()))
        self.errors = errors


class AlreadyPrinted(ProductionError):
    pass


class PossibleDuplicate(ProductionError):
    """The same document line seems to have a label already (this or another computer)."""

    def __init__(self, barcodes: list[str], ref: str):
        super().__init__(f"line {ref} already has label(s) {', '.join(barcodes)}")
        self.barcodes = barcodes
        self.ref = ref


class PrintInProgress(ProductionError):
    pass


class StationNotVerified(ProductionError):
    pass


class PreviousPrintUnknown(ProductionError):
    """An earlier print request for this label got no answer: it may or may not have printed."""

    def __init__(self, barcode: str, job_keys: list[str]):
        super().__init__(f"An earlier print request for {barcode} got no answer from the browser. Check the "
                         "printer: did the label come out?")
        self.barcode = barcode
        self.job_keys = job_keys


PRINT_LOCK_SECONDS = 120


@dataclass
class Outcome:
    label_id: Optional[int]
    barcode: Optional[str]
    print_status: str
    excel_status: str
    messages: list[str] = field(default_factory=list)
    print_ok: bool = False


class ProductionService:
    def __init__(self, *, settings: Settings, db: Database, repo: Repository, templates: TemplateRegistry,
                 transforms: Transforms, allocator: BarcodeAllocator, printer: LabelPrinter,
                 writer: SafeWorkbookWriter, registry_client: Optional[RegistryAllocator] = None,
                 clock=None):
        self.s = settings
        self.db = db
        self.repo = repo
        self.templates = templates
        self.transforms = transforms
        self.allocator = allocator
        self.printer = printer
        self.writer = writer
        self.registry_client = registry_client
        self.tz = ZoneInfo(settings.timezone)
        self._clock = clock or (lambda: datetime.now(timezone.utc))
        self.labels_dir = settings.data_path / "labels"
        self.browser = None
        self.cloud_excel = None  # WorkbookSync (VSPM_EXCEL_MODE=onedrive), set by build_context

    # ---- helpers ------------------------------------------------------------------
    def _m(self, key: str, **kw) -> str:
        el, en = MESSAGES[key]
        return (el if self.s.language == "el" else en).format(**kw)

    def now_local(self) -> datetime:
        return self._clock().astimezone(self.tz)

    def install_id(self) -> str:
        v = self.db.get_setting("install_id")
        if not v:
            v = uuid.uuid4().hex[:12]
            self.db.set_setting("install_id", v)
        return v

    def barcode_key(self, row_id: int) -> str:
        return f"vspm:{self.install_id()}:row:{row_id}"

    def _who(self, operator: Optional[str]) -> tuple[str, str]:
        return operator or self.s.effective_operator, self.s.effective_computer_name

    def check_printable(self, row: ProductionRow) -> dict[str, str]:
        errors = validate_row(row)
        if not row.approved:
            errors.setdefault("approved", "row must be approved by the operator")
        if row.template_key is not None:
            try:
                spec = self.templates.spec(row.template_key)
                self.templates.load(row.template_key)
                build_label_content(row, spec, self.transforms, barcode="X")
            except TransformError as e:
                errors["material"] = str(e)
            except Exception as e:  # missing / changed template file
                errors["template_key"] = str(e)
        return errors

    def preview(self, row: ProductionRow, barcode: Optional[str] = None) -> tuple[bytes, str, LabelContent]:
        """PNG preview for a row (proposed or existing barcode)."""
        if row.template_key is None:
            raise ProductionError("Select a label template first")
        spec = self.templates.spec(row.template_key)
        tpl = self.templates.load(row.template_key)
        label = self.repo.label_for_row(row.id) if row.id else None
        if label is not None:
            content = LabelContent.model_validate_json(label["content_json"])
            number = label["number"]
        else:
            ci = tpl.counter_info(spec.counter_object) or {}
            placeholder = barcode or f"{spec.barcode_prefix}{'0' * max(ci.get('format_width', 3), 3)}"
            content = build_label_content(row, spec, self.transforms, placeholder)
            number = 0
        xml = tpl.populate(spec, content, number)
        png = self.printer.render_preview(xml) or render_label_png(xml, now=self.now_local().replace(tzinfo=None),
                                                                    lang=self.s.language)
        return png, xml, content

    # ---- Print & Save ---------------------------------------------------------------
    def production_ref(self, row: ProductionRow) -> str:
        """Identity of the paper line: same scan + same row number = same production line,
        on every computer."""
        doc = self.repo.get_document(row.document_id) if row.document_id else None
        if doc is None:
            return f"local:{self.install_id()}:row{row.id}"
        return f"{doc['source']}:{doc['source_file_id']}#row{row.row_index}"

    def find_duplicates(self, row: ProductionRow) -> list[str]:
        """Barcodes already issued for the same document line (excluding this row's own)."""
        ref = self.production_ref(row)
        own_key = self.barcode_key(row.id)
        found = [r["barcode"] for r in self.db.query(
            "SELECT barcode FROM labels WHERE production_ref=? AND production_row_id<>?", (ref, row.id))]
        finder = getattr(self.allocator, "find_by_ref", None)
        if finder is not None:
            try:
                for res in finder(ref):
                    if res.idempotency_key != own_key and res.status != "void" and res.barcode not in found:
                        found.append(res.barcode)
            except BarcodeError as e:
                log.warning("duplicate check against registry failed: %s", e)
        return found

    def print_and_save(self, row_id: int, operator: Optional[str] = None,
                       allow_possible_duplicate: bool = False) -> Outcome:
        op, _ = self._who(operator)
        label = self.prepare_label(row_id, op, allow_possible_duplicate)
        # Resume (or first attempt) of printing with the reserved number.
        return self._print(label, kind="initial", operator=op)

    def prepare_label(self, row_id: int, operator: Optional[str] = None,
                      allow_possible_duplicate: bool = False):
        """Validate the row, reserve its barcode and COMMIT the label row (idempotent).

        Returns the label row. Raises AlreadyPrinted if the row was printed before."""
        op, pc = self._who(operator)
        row = self.repo.get_row(row_id)
        if row is None:
            raise ProductionError(f"Row {row_id} not found")
        label = self.repo.label_for_row(row_id)
        if label is not None and label["print_status"] in (
                PrintStatus.print_submitted.value, PrintStatus.print_unconfirmed.value,
                PrintStatus.printed_confirmed.value, PrintStatus.print_failed_confirmed.value):
            raise AlreadyPrinted(f"Row already printed as {label['barcode']}. Use 'Reprint' instead.")
        if label is None:
            errors = self.check_printable(row)
            if errors:
                raise ValidationFailed(errors)
            dupes = self.find_duplicates(row)
            if dupes and not allow_possible_duplicate:
                raise PossibleDuplicate(dupes, self.production_ref(row))
            if dupes:
                self.db.audit("row", row.id, "possible_duplicate_overridden", {"barcodes": dupes}, op, pc)
            spec = self.templates.spec(row.template_key)
            tpl = self.templates.load(row.template_key)
            res = self._reserve(row, spec.barcode_prefix, op, pc)
            content = build_label_content(row, spec, self.transforms, res.barcode)
            label_id = self._create_label(row, res, content, op, pc, template_sha256=tpl.sha256)
            label = self.repo.get_label(label_id)
        return label

    def _reserve(self, row: ProductionRow, series: str, op: str, pc: str) -> Reservation:
        key = self.barcode_key(row.id)
        try:
            res = self.allocator.reserve(series, key, production_ref=self.production_ref(row), operator=op,
                                         computer=pc)
        except BarcodeError as e:
            self.db.audit("row", row.id, "barcode_reservation_failed", {"error": str(e)}, op, pc)
            raise ProductionError(str(e)) from e
        if res.series != series:
            raise ProductionError(f"Registry returned series {res.series}, expected {series}")
        return res

    def _create_label(self, row: ProductionRow, res: Reservation, content: LabelContent, op: str, pc: str,
                      template_sha256: Optional[str] = None) -> int:
        now = utcnow()
        with self.db.transaction() as conn:
            existing = conn.execute("SELECT id FROM labels WHERE production_row_id=?", (row.id,)).fetchone()
            if existing:
                return existing["id"]
            lid = conn.execute(
                "INSERT INTO labels (production_row_id, template_key, series, number, barcode, reservation_id, "
                "reservation_mode, idempotency_key, content_json, print_status, excel_status, operator, computer, "
                "created_at, updated_at, production_ref, template_sha256) VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
                (row.id, row.template_key.value, res.series, res.number, res.barcode, res.reservation_id, res.mode,
                 res.idempotency_key, content.model_dump_json(), PrintStatus.barcode_reserved.value,
                 ExcelStatus.not_queued.value, op, pc, now, now, self.production_ref(row), template_sha256)).lastrowid
            self.repo.set_row_status(row.id, RowStatus.barcode_reserved, conn)
            self.db.audit("label", lid, "barcode_reserved",
                          {"barcode": res.barcode, "row": row.id, "mode": res.mode}, op, pc, conn=conn)
        return lid

    def _label_xml(self, label) -> str:
        key = TemplateKey(label["template_key"])
        spec = self.templates.spec(key)
        content = LabelContent.model_validate_json(label["content_json"])
        tpl = self.templates.load_version(key, label["template_sha256"]) or self.templates.load(key)
        return tpl.populate(spec, content, label["number"])

    def _claim_print_lock(self, label_id: int) -> bool:
        """Only one print job per label at a time (e.g. two browser tabs)."""
        now = self._clock()
        until = (now + timedelta(seconds=PRINT_LOCK_SECONDS)).isoformat()
        with self.db.transaction() as conn:
            cur = conn.execute("UPDATE labels SET print_lock_until=? WHERE id=? AND "
                               "(print_lock_until IS NULL OR print_lock_until < ?)", (until, label_id, now.isoformat()))
            return cur.rowcount == 1

    def _release_print_lock(self, label_id: int) -> None:
        self.db.execute("UPDATE labels SET print_lock_until=NULL WHERE id=?", (label_id,))

    def _print(self, label, kind: str, operator: str) -> Outcome:
        if not self._claim_print_lock(label["id"]):
            raise PrintInProgress(f"{label['barcode']} is being printed right now (another window/tab). Wait a moment.")
        try:
            # Re-read: another session may have printed it while we waited for the lock.
            fresh = self.repo.get_label(label["id"])
            if kind == "initial" and fresh["first_printed_at"]:
                raise AlreadyPrinted(f"Row already printed as {fresh['barcode']}. Use 'Reprint' instead.")
            return self._print_locked(fresh, kind, operator)
        finally:
            self._release_print_lock(label["id"])

    def _print_locked(self, label, kind: str, operator: str) -> Outcome:
        op, pc = self._who(operator)
        xml = self._label_xml(label)
        self.labels_dir.mkdir(parents=True, exist_ok=True)
        xml_path = self.labels_dir / f"{label['barcode']}.label"
        xml_path.write_text(xml, encoding="utf-8")
        result = self.printer.print_label(xml, self.s.printer_name)
        return self.record_print_outcome(label, kind, result.ok, result.error, result.job_ref, self.printer.mode,
                                         self.s.printer_name, op, pc, xml_path=str(xml_path))

    def record_print_outcome(self, label, kind: str, ok: bool, error: Optional[str], job_ref: Optional[str],
                             printer_mode: str, printer_name: Optional[str], operator: str, computer: str,
                             xml_path: Optional[str] = None, confirmed: bool = False) -> Outcome:
        """Record one print submission (server printer, browser or manual DYMO Label).

        The first successful submission sets the print date and queues the Excel row;
        reprints never touch Excel."""
        op, pc = operator, computer
        lid = label["id"]
        label = self.repo.get_label(lid)
        self.repo.add_print_attempt(lid, kind=kind, printer_mode=printer_mode, printer_name=printer_name or "",
                                    status=("confirmed_ok" if confirmed else "submitted") if ok else "error",
                                    error=error, job_ref=job_ref, operator=op, computer=pc)
        msgs: list[str] = []
        extra = {"label_xml_path": xml_path} if xml_path else {}
        if not ok:
            if label["first_printed_at"]:
                status = label["print_status"]  # an earlier successful print stays valid
            else:
                status = PrintStatus.error.value
                self.repo.set_row_status(label["production_row_id"], RowStatus.error)
            self.repo.update_label(lid, print_status=status, **extra)
            self.db.audit("label", lid, f"{kind}_failed", {"error": error, "job": job_ref}, op, pc)
            msgs.append(self._m("print_failed", e=error))
            lab = self.repo.get_label(lid)
            return Outcome(lid, lab["barcode"], lab["print_status"], lab["excel_status"], msgs, False)

        first_success = not label["first_printed_at"]
        new_status = PrintStatus.printed_confirmed if confirmed else PrintStatus.print_unconfirmed
        fields = dict(print_status=new_status.value, **extra)
        if first_success:
            local = self.now_local()
            fields.update(first_printed_at=self._clock().isoformat(timespec="seconds"),
                          print_date_local=local.date().isoformat())
        self.repo.update_label(lid, **fields)
        self.repo.set_row_status(label["production_row_id"],
                                 RowStatus.printed_confirmed if confirmed else RowStatus.print_unconfirmed)
        self.db.audit("label", lid, f"{kind}_submitted", {"job": job_ref, "barcode": label["barcode"],
                                                          "printer": printer_name, "mode": printer_mode}, op, pc)
        msgs.append(self._m("sent") if not confirmed else self._m("manual_recorded"))
        if first_success:
            try:
                self.allocator.mark(Reservation(label["series"], label["number"], label["barcode"],
                                                label["reservation_id"] or "", label["idempotency_key"],
                                                label["reservation_mode"]), "printed", op)
            except Exception as e:  # registry bookkeeping only
                log.warning("Could not mark reservation printed: %s", e)
        lab = self.repo.get_label(lid)
        if lab["excel_status"] == ExcelStatus.not_queued.value:
            self._queue_excel(lab, op)
            msgs.extend(self.sync_excel(lid, operator=op))
        lab = self.repo.get_label(lid)
        return Outcome(lid, lab["barcode"], lab["print_status"], lab["excel_status"], msgs, True)

    # ---- browser printing (the employee's own PC prints) ----------------------------------
    def _check_station(self, station_id: Optional[str], printer: Optional[str]) -> None:
        if self.browser is None:
            raise ProductionError("Browser printing is not configured")
        if self.s.require_verified_station and not self.browser.station_verified(station_id, printer):
            raise StationNotVerified(
                "This computer/printer has not printed TEST0001 successfully yet. Open the 'DYMO printer' page, "
                "print TEST0001 and confirm it before printing production labels.")

    def start_browser_print(self, row_id: int, *, session_id: str, station_id: Optional[str],
                            printer: Optional[str], operator: Optional[str] = None,
                            allow_possible_duplicate: bool = False, previous_not_printed: bool = False):
        """Reserve + persist the barcode, then create a print job for THIS browser."""
        op, _ = self._who(operator)
        self._check_station(station_id, printer)  # before any number is reserved
        label = self.prepare_label(row_id, op, allow_possible_duplicate)
        unknown = self.browser.unanswered_jobs_for_label(label["id"])
        if unknown and not previous_not_printed:
            raise PreviousPrintUnknown(label["barcode"], [j["job_key"] for j in unknown])
        if unknown:
            self._resolve_unknown(label, unknown, printed=False, operator=op)
        kind = "retry" if self.repo.print_attempts(label["id"]) else "initial"
        content = LabelContent.model_validate_json(label["content_json"])
        return self.browser.create_label_job(label, kind, session_id, op, content)

    def start_browser_reprint(self, label_id: int, *, session_id: str, station_id: Optional[str],
                              printer: Optional[str], operator: Optional[str] = None):
        op, _ = self._who(operator)
        self._check_station(station_id, printer)
        label = self.repo.get_label(label_id)
        if label is None:
            raise ProductionError(f"Label {label_id} not found")
        if not label["first_printed_at"]:
            raise ProductionError("This label was never printed: use Print & Save")
        content = LabelContent.model_validate_json(label["content_json"])
        return self.browser.create_label_job(label, "reprint", session_id, op, content)

    def apply_browser_result(self, job, result: dict, user: str) -> list[tuple[str, str]]:
        """Callback from BrowserPrintService for production jobs (idempotent per job)."""
        label = self.repo.get_label(job["label_id"])
        status = result.get("status")
        computer = result.get("station_name") or result.get("station_id") or "browser"
        if status == "submitted":
            out = self.record_print_outcome(label, job["kind"], True, None, job["job_key"], "browser",
                                            result.get("printer"), user, computer)
            return [("success", f"{label['barcode']}: {m}") for m in out.messages]
        if status in ("failed", "rejected"):
            err = result.get("error") or "; ".join(result.get("problems") or []) or status
            out = self.record_print_outcome(label, job["kind"], False, err, job["job_key"], "browser",
                                            result.get("printer"), user, computer)
            return [("error", f"{label['barcode']}: {m}") for m in out.messages]
        if status == "uncertain":
            return [("warning", self._m("print_uncertain", b=label["barcode"]))]
        return [("warning", f"{label['barcode']}: {result.get('error') or status}")]  # not_submitted

    def _resolve_unknown(self, label, jobs, printed: bool, operator: str) -> None:
        for j in jobs:
            self.db.execute("UPDATE print_jobs SET status=? WHERE job_key=?",
                            ("manual_confirmed" if printed else "cancelled", j["job_key"]))
            self.db.audit("label", label["id"], "unanswered_print_resolved",
                          {"job": j["job_key"], "printed": printed}, operator)

    def confirm_unanswered_printed(self, label_id: int, operator: Optional[str] = None) -> Outcome:
        """The operator checked the printer: the label of an unanswered job DID come out."""
        op, pc = self._who(operator)
        label = self.repo.get_label(label_id)
        unknown = self.browser.unanswered_jobs_for_label(label_id)
        if not unknown:
            raise ProductionError("No unanswered print request for this label")
        self._resolve_unknown(label, unknown, printed=True, operator=op)
        return self.record_print_outcome(label, unknown[-1]["kind"], True, None, unknown[-1]["job_key"],
                                         "browser", unknown[-1]["printer_name"], op, pc, confirmed=True)

    # ---- manual fallback: download the prepared .label and print it in DYMO Label v8 ---------
    def manual_label_file(self, label_id: int) -> tuple[str, bytes]:
        """(file name, bytes) of the exact prepared label with the reserved barcode (nothing recorded)."""
        if self.browser is None:
            raise ProductionError("Browser printing is not configured")
        label = self.repo.get_label(label_id)
        content = LabelContent.model_validate_json(label["content_json"])
        job = self.browser.prepare("manual", TemplateKey(label["template_key"]), content, label["number"],
                                   label["id"], "", template_sha256=label["template_sha256"])
        return f"{label['barcode']}.label", job.label_xml.encode("utf-8")

    def record_manual_download(self, label_id: int, operator: Optional[str] = None) -> str:
        """The operator downloaded the .label file to print it in DYMO Label v8."""
        op, _ = self._who(operator)
        existing = self.db.one("SELECT job_key FROM print_jobs WHERE label_id=? AND status='manual_pending'",
                               (label_id,))
        if existing is not None:
            return existing["job_key"]
        label = self.repo.get_label(label_id)
        _, data = self.manual_label_file(label_id)
        from vspm.dymo.template import xml_sha256
        key = uuid.uuid4().hex
        with self.db.transaction() as conn:
            conn.execute(
                "INSERT INTO print_jobs (job_key, kind, label_id, template_key, barcode, xml_sha256, status, "
                "created_by, created_at) VALUES (?,?,?,?,?,?, 'manual_pending', ?, ?)",
                (key, "manual", label_id, label["template_key"], label["barcode"], xml_sha256(data.decode("utf-8")),
                 op, utcnow()))
            self.db.audit("label", label_id, "manual_label_downloaded", {"job": key}, op, conn=conn)
        return key

    def confirm_manual_print(self, label_id: int, printed_ok: bool, operator: Optional[str] = None) -> Outcome:
        """Operator printed the downloaded .label in DYMO Label v8 and checked the result."""
        op, pc = self._who(operator)
        label = self.repo.get_label(label_id)
        job = self.db.one("SELECT * FROM print_jobs WHERE label_id=? AND status='manual_pending'", (label_id,))
        if job is None:
            raise ProductionError("Download the .label file first")
        self.db.execute("UPDATE print_jobs SET status=?, reported_at=? WHERE job_key=?",
                        ("manual_confirmed" if printed_ok else "cancelled", utcnow(), job["job_key"]))
        kind = "manual_reprint" if label["first_printed_at"] else "manual"
        if not printed_ok:
            return self.record_print_outcome(label, kind, False, "not printed (operator)", job["job_key"],
                                             "manual_dymo_label", None, op, pc)
        return self.record_print_outcome(label, kind, True, None, job["job_key"], "manual_dymo_label", None, op, pc,
                                         confirmed=True)

    def reprint(self, label_id: int, operator: Optional[str] = None) -> Outcome:
        label = self.repo.get_label(label_id)
        if label is None:
            raise ProductionError(f"Label {label_id} not found")
        kind = "reprint" if label["first_printed_at"] else "retry"
        return self._print(label, kind=kind, operator=operator or self.s.effective_operator)

    def confirm_print(self, label_id: int, ok: bool, operator: Optional[str] = None) -> None:
        op, pc = self._who(operator)
        label = self.repo.get_label(label_id)
        if label is None or not label["first_printed_at"]:
            raise ProductionError("Nothing to confirm: the label was never sent to the printer")
        status = PrintStatus.printed_confirmed if ok else PrintStatus.print_failed_confirmed
        self.repo.update_label(label_id, print_status=status.value)
        self.repo.set_row_status(label["production_row_id"],
                                 RowStatus.printed_confirmed if ok else RowStatus.print_unconfirmed)
        last = self.repo.print_attempts(label_id)[-1]
        self.db.execute("UPDATE print_attempts SET status=?, confirmed_at=?, confirmed_by=? WHERE id=?",
                        ("confirmed_ok" if ok else "confirmed_failed", utcnow(), op, last["id"]))
        self.db.audit("label", label_id, "print_confirmed" if ok else "print_reported_failed", {}, op, pc)

    # ---- Excel --------------------------------------------------------------------------
    def _payload(self, label) -> ExcelRowPayload:
        if label["excel_payload_json"]:
            return ExcelRowPayload.model_validate_json(label["excel_payload_json"])
        row = self.repo.get_row(label["production_row_id"])
        content = LabelContent.model_validate_json(label["content_json"])
        spec = self.templates.spec(TemplateKey(label["template_key"]))
        return build_excel_payload(row, content, spec, self.transforms, date.fromisoformat(label["print_date_local"]),
                                   self.s.excel_material_source)

    def _queue_excel(self, label, op: str) -> None:
        payload = self._payload(label)
        self.repo.update_label(label["id"], excel_status=ExcelStatus.pending.value,
                               excel_payload_json=payload.model_dump_json())
        self.db.audit("label", label["id"], "excel_queued", payload.model_dump(), op, self.s.effective_computer_name)

    def sync_excel(self, label_id: int, operator: Optional[str] = None) -> list[str]:
        """Try to record the label in Excel according to the configured mode."""
        op, pc = self._who(operator)
        label = self.repo.get_label(label_id)
        if label is None or label["excel_status"] == ExcelStatus.not_queued.value:
            return [self._m("excel_nothing")]
        if label["excel_status"] == ExcelStatus.synchronized.value:
            return [self._m("excel_already")]
        payload = self._payload(label)
        mode = self.s.excel_mode
        if mode == ExcelMode.staging:
            self.repo.update_label(label_id, excel_status=ExcelStatus.staged.value)
            self.repo.add_excel_attempt(label_id, mode=mode.value, status="staged")
            return [self._m("excel_staged")]
        if mode == ExcelMode.onedrive:
            # The row is already stored in the database; the single OneDrive writer adds it to the workbook
            # (or keeps it staged until automatic writing has been tested on a copy and enabled).
            if self.cloud_excel is None:
                return self._excel_failed(label_id, mode.value, "OneDrive workbook sync is not configured", "error")
            res = self.cloud_excel.sync_pending(op, label_ids=[label_id])
            if res.status == "synchronized":
                return [self._m("excel_written", n=res.rows.get(payload.barcode, "?"))]
            if res.status == "staged":
                return [self._m("excel_staged")]
            return [f"Excel (OneDrive): {res.message}"]
        if mode == ExcelMode.registry_queue:
            if self.registry_client is None:
                return self._excel_failed(label_id, mode.value, "registry_queue mode needs VSPM_REGISTRY_URL/API key",
                                          "error")
            try:
                self.registry_client.enqueue_excel(f"excel:{payload.barcode}", payload.barcode, payload.model_dump())
            except BarcodeError as e:
                return self._excel_failed(label_id, mode.value, str(e), "error")
            self.repo.update_label(label_id, excel_status=ExcelStatus.queued_remote.value)
            self.repo.add_excel_attempt(label_id, mode=mode.value, status="queued_remote")
            return [self._m("excel_queued")]
        return self._direct_write(label_id, [payload], op)[label_id]

    def _direct_write(self, label_id: int, payloads: list[ExcelRowPayload], op: str) -> dict[int, list[str]]:
        try:
            report = self.writer.append(payloads)
        except ExcelSyncError as e:
            return {label_id: self._excel_failed(label_id, "direct", str(e), e.status)}
        except Exception as e:  # unexpected - keep the label recoverable
            log.exception("Excel write failed")
            return {label_id: self._excel_failed(label_id, "direct", f"Unexpected error: {e}", "error")}
        res = report.results[0]
        self.repo.update_label(label_id, excel_status=ExcelStatus.synchronized.value, excel_row_number=res.row_number)
        self.repo.add_excel_attempt(label_id, mode="direct", status="synchronized" if res.status == "written"
                                    else "already_present", row_number=res.row_number,
                                    workbook_path=str(self.writer.path), hash_before=report.hash_before,
                                    hash_after=report.hash_after,
                                    backup_path=str(report.backup_path) if report.backup_path else None)
        self.db.audit("label", label_id, "excel_synchronized", {"row": res.row_number, "result": res.status}, op,
                      self.s.effective_computer_name)
        key = "excel_written" if res.status == "written" else "excel_present"
        return {label_id: [self._m(key, n=res.row_number)]}

    def _excel_failed(self, label_id: int, mode: str, error: str, status: str) -> list[str]:
        st = ExcelStatus.conflict if status == "conflict" else ExcelStatus.error
        self.repo.update_label(label_id, excel_status=st.value)
        self.repo.add_excel_attempt(label_id, mode=mode, status=st.value, error=error,
                                    workbook_path=str(self.writer.path))
        self.db.audit("label", label_id, "excel_failed", {"error": error, "mode": mode})
        return [self._m("excel_failed", st=st.value, e=error)]

    def apply_staged(self, label_ids: Optional[list[int]] = None, operator: Optional[str] = None) -> dict[int, list[str]]:
        """Write staged/pending/failed rows directly (operator-initiated)."""
        op, _ = self._who(operator)
        rows = self.pending_excel_labels() if label_ids is None else [self.repo.get_label(i) for i in label_ids]
        out: dict[int, list[str]] = {}
        for lab in rows:
            if lab["excel_status"] in (ExcelStatus.synchronized.value, ExcelStatus.not_queued.value):
                continue
            out.update(self._direct_write(lab["id"], [self._payload(lab)], op))
        return out

    def pending_excel_labels(self):
        return self.db.query("SELECT * FROM labels WHERE excel_status IN ('pending','staged','error','conflict') "
                             "ORDER BY id")

    def reconcile(self, workbook: bytes | None = None, source: str | None = None) -> dict[str, list[str]]:
        """Read the workbook (read-only) and mark staged/failed labels that are present as synchronized.

        `workbook`: bytes of an uploaded copy (cloud: the server cannot open the OneDrive file)."""
        labs = self.db.query("SELECT * FROM labels WHERE excel_status IN ('pending','staged','error','conflict',"
                             "'queued_remote')")
        if not labs:
            return {"synchronized": [], "missing": [], "duplicates": []}
        try:
            found = self.writer.find_barcodes([l["barcode"] for l in labs], workbook)
        except Exception as e:
            raise ProductionError(f"Workbook could not be read: {e}") from e
        where = source or str(self.writer.path)
        out = {"synchronized": [], "missing": [], "duplicates": []}
        for l in labs:
            rows = found.get(l["barcode"], [])
            if len(rows) == 1:
                self.repo.update_label(l["id"], excel_status=ExcelStatus.synchronized.value, excel_row_number=rows[0])
                self.repo.add_excel_attempt(l["id"], mode="reconcile", status="already_present", row_number=rows[0],
                                            workbook_path=where)
                out["synchronized"].append(l["barcode"])
            elif len(rows) > 1:
                self.repo.update_label(l["id"], excel_status=ExcelStatus.conflict.value)
                self.repo.add_excel_attempt(l["id"], mode="reconcile", status="conflict",
                                            error=f"barcode appears in rows {rows}", workbook_path=where)
                out["duplicates"].append(f"{l['barcode']} rows {rows}")
            else:
                out["missing"].append(l["barcode"])
        return out

    def refresh_remote(self) -> int:
        if self.registry_client is None:
            return 0
        n = 0
        for l in self.db.query("SELECT * FROM labels WHERE excel_status='queued_remote'"):
            job = self.registry_client.excel_job(f"excel:{l['barcode']}")
            if job and job["status"] in ("done", "already_present"):
                result = json.loads(job.get("result") or "{}")
                self.repo.update_label(l["id"], excel_status=ExcelStatus.synchronized.value,
                                       excel_row_number=result.get("row_number"))
                n += 1
            elif job and job["status"] in ("error", "conflict"):
                self.repo.update_label(l["id"], excel_status=job["status"])
        return n

    def export_staging(self, path: Path) -> int:
        """Values-only sheet of rows not yet in Excel, in the workbook's column order."""
        import openpyxl
        labs = self.db.query("SELECT * FROM labels WHERE excel_status IN ('pending','staged','error','conflict') "
                             "ORDER BY id")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "ΓΙΑ ΚΑΤΑΧΩΡΗΣΗ"
        ws.append(["ΕΙΔΟΣ", "ΠΑΧΟΣ", "ΔΙΑΣΤΑΣΕΙΣ", "ΤΕΜΑΧΙΑ", "ΚΙΛΑ", "ΠΕΛΑΤΗΣ", "VS/PG", "BARCODE",
                   "ΤΟΠΟΘΕΣΙΑ", "ΘΕΣΗ", "ΡΟΛΟ", "ΠΟΙΟΤΗΤΑ", "ΧΩΡΑ", "ID", "ΚΟΣΤΟΣ", "ΗΜΕΡΟΜΗΝΙΑ"])
        from vspm.excel.writer import payload_to_cells
        for l in labs:
            v, d = payload_to_cells(self._payload(l))
            ws.append([v["A"], v["B"], v["C"], v["D"], v["E"], v["F"], v["G"], v["H"], v["I"], None, v["K"],
                       None, None, None, None, d["P"]])
            ws.cell(ws.max_row, 16).number_format = "dd/mm/yyyy"
        path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(path)
        return len(labs)

    # ---- recovery -----------------------------------------------------------------------------
    def recover(self) -> list[str]:
        """Called at start-up: finish steps interrupted by a crash/restart."""
        notes = []
        for l in self.db.query("SELECT * FROM labels WHERE first_printed_at IS NOT NULL AND excel_status='not_queued'"):
            self._queue_excel(l, "recovery")
            notes.append(self._m("rec_queued", b=l["barcode"]))
        for l in self.db.query("SELECT * FROM labels WHERE print_status='barcode_reserved'"):
            notes.append(self._m("rec_reserved", b=l["barcode"]))
        return notes
