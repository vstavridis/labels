"""Synthetic workbook for tests and dry runs (same structure as the production sheet,
no real data)."""

from __future__ import annotations

import io
import zipfile
from datetime import datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

from vspm.config import DEFAULT_SHEET_NAME

FORMULAS = {
    "L": '=IFERROR(INDEX([1]!Table4[ΠΟΙΟΤΗΤΑ],MATCH(K{r},[1]!Table4[Νο ΡΟΛΛΟΥ],0)),"-")',
    "M": '=IFERROR(VLOOKUP(K{r},[1]ΣΒΗΣΜΕΝΑ!$F:$G,2,),"-")',
    "N": '=IFERROR(VLOOKUP(F{r},ID!$A:$B,2,0),"-")',
    "O": '=IFERROR(VLOOKUP(K{r},[1]ΣΒΗΣΜΕΝΑ!$F:$J,5,),"-")',
}
EXTRA_PARTS = {
    "xl/ctrlProps/ctrlProp1.xml": b'<?xml version="1.0"?><formControlPr objectType="Button"/>',
    "xl/custom/keepme.bin": bytes(range(256)),
}


def build_test_workbook(path: Path, data_rows: int = 3, styled_empty_rows: int = 5, hidden_after: bool = True) -> Path:
    wb = openpyxl.Workbook()
    ids = wb.active
    ids.title = "ID"
    ids.append(["Customer", 5])
    ids.append(["TEST CUSTOMER", 9])
    ws = wb.create_sheet(DEFAULT_SHEET_NAME)
    ws.append(["ΕΙΔΟΣ", "ΠΑΧΟΣ", "ΔΙΑΣΤΑΣΕΙΣ", "ΤΕΜΑΧΙΑ", "ΚΙΛΑ", "ΠΕΛΑΤΗΣ", "VS/PG", "BARCODE", "ΤΟΠΟΘΕΣΙΑ",
               "ΘΕΣΗ", "ΡΟΛΟ", "ΠΟΙΟΤΗΤΑ", "ΧΩΡΑ", "ID", "ΚΟΣΤΟΣ", "ΗΜΕΡΟΜΗΝΙΑ"])
    for c in range(1, 17):
        ws.cell(1, c).font = Font(bold=True)
    for i in range(data_rows):
        r = i + 2
        ws.append(["ΓΑΛΒΑΝΙΖΕ", "0.60", f"{30 + i}ΧΡΟΛΟ", 2, 200 + i, "ΑΠΟΘΗΚΗ", "VS", f"VS{9000 + i}", "ΜΙΚΡΟ",
                   None, f"ROLL{i}"] + [FORMULAS[c].format(r=r) for c in "LMNO"] + [datetime(2025, 1, 10 + i)])
        ws.cell(r, 16).number_format = "mm-dd-yy"
        ws.cell(r, 10).value = "Α" if i == 0 else None
    for j in range(styled_empty_rows):
        r = data_rows + 2 + j
        for c in range(1, 17):
            ws.cell(r, c).font = Font(italic=False, size=11)
        ws.cell(r, 16).number_format = "mm-dd-yy"
        if hidden_after and j >= 2:
            ws.row_dimensions[r].hidden = True
    wb.create_sheet("Σβησμενα ")
    buf = io.BytesIO()
    wb.save(buf)
    # Add parts openpyxl would not preserve, to prove we keep them untouched.
    src = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            z.writestr(info, src.read(info.filename))
        for name, data in EXTRA_PARTS.items():
            z.writestr(name, data)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(out.getvalue())
    return path


