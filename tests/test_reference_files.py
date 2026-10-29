"""Checks against the real reference files. Skipped when they are not present
(they are confidential and not committed). Never touches the originals: the
workbook is copied to a temporary folder first."""

from __future__ import annotations

import hashlib
import shutil
import subprocess
import zipfile

import openpyxl
import pytest

from vspm.config import DEFAULT_SHEET_NAME
from vspm.domain.models import ExcelRowPayload, TemplateKey
from vspm.dymo.template import TemplateRegistry
from vspm.excel.formulas import row_pattern
from vspm.excel.writer import SafeWorkbookWriter
from vspm.excel.xlsx_package import SheetEditor, XlsxPackage
from tests.conftest import REFERENCE, ROOT

WORKBOOK = REFERENCE / "Σχισμενα.xlsx"
pytestmark = pytest.mark.skipif(not WORKBOOK.exists(), reason="reference files not available")


def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()


def test_real_templates_match_configuration():
    reg = TemplateRegistry.from_yaml(ROOT / "config" / "templates.yaml", REFERENCE)
    for key in TemplateKey:
        tpl = reg.load(key)  # raises if structure differs from templates.yaml
        bc = tpl.find("BARCODE")
        assert bc.findtext("Type") == "Code39" and bc.findtext("LinkedObjectName") == "COUNTER"


def test_real_sheet_name_and_formulas():
    ed = SheetEditor(XlsxPackage(WORKBOOK.read_bytes()), DEFAULT_SHEET_NAME)
    assert DEFAULT_SHEET_NAME[0] == "΄"
    last = ed.last_data_row()
    for col in "LMNO":
        pats = {row_pattern(ed.value(r, col).formula, r) for r in range(2, last + 1)}
        assert len(pats) == 1, f"column {col} has several formula patterns"


def test_write_to_copy_of_real_workbook(tmp_path):
    original_hash = sha(WORKBOOK)
    copy = tmp_path / "Σχισμενα.xlsx"
    shutil.copyfile(WORKBOOK, copy)
    before = zipfile.ZipFile(copy)
    parts = {n: before.read(n) for n in before.namelist()}
    w = SafeWorkbookWriter(copy, DEFAULT_SHEET_NAME, backup_dir=tmp_path / "bk", write_enabled=True)
    last = SheetEditor(XlsxPackage(copy.read_bytes()), DEFAULT_SHEET_NAME).last_data_row()
    p = ExcelRowPayload(material="ΓΑΛΒΑΝΙΖΕ", thickness="0.60", dimensions="34ΧΡΟΛΟ", pieces="4", weight="496",
                        customer="ΤΕΣΤ", category="VS", barcode="VSTEST1", location="7", roll="TESTROLL",
                        print_date="2026-10-21")
    rep = w.append([p])
    assert rep.results[0].row_number == last + 1
    after = zipfile.ZipFile(copy)
    changed = {n for n in parts if after.read(n) != parts[n]}
    assert changed <= {"xl/workbook.xml", "xl/styles.xml", "xl/worksheets/sheet6.xml"}
    ed = SheetEditor(XlsxPackage(copy.read_bytes()), DEFAULT_SHEET_NAME)
    r = last + 1
    for col in "LMNO":
        assert row_pattern(ed.value(r, col).formula, r) == row_pattern(ed.value(last, col).formula, last)
        assert "[1]" in ed.value(r, col).formula or col == "N"
    assert sha(WORKBOOK) == original_hash  # original untouched


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
def test_written_copy_opens_in_libreoffice(tmp_path):
    copy = tmp_path / "wb.xlsx"
    shutil.copyfile(WORKBOOK, copy)
    w = SafeWorkbookWriter(copy, DEFAULT_SHEET_NAME, backup_dir=tmp_path / "bk", write_enabled=True)
    w.append([ExcelRowPayload(material="ΓΑΛΒΑΝΙΖΕ", thickness="0.60", dimensions="34ΧΡΟΛΟ", pieces="4", weight="496",
                              customer="ΤΕΣΤ", category="VS", barcode="VSTEST2", location="7", roll="X",
                              print_date="2026-10-21")])
    out = tmp_path / "out"
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(out), str(copy)],
                   check=True, capture_output=True, timeout=240)
    ws = openpyxl.load_workbook(out / "wb.xlsx", data_only=True)[DEFAULT_SHEET_NAME]
    rows = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 8).value == "VSTEST2"]
    assert len(rows) == 1
    assert ws.cell(rows[0], 16).value.date().isoformat() == "2026-10-21"
