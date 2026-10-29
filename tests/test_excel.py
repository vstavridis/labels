"""Excel column mapping, formulas, dates, idempotency and safety checks."""

from __future__ import annotations

import zipfile
from datetime import date, datetime

import openpyxl
import pytest

from vspm.config import DEFAULT_SHEET_NAME
from vspm.domain.models import ExcelRowPayload
from vspm.excel.formulas import FALLBACK_FORMULAS, row_pattern, shift_rows, translate_to_row
from vspm.excel.writer import (
    ExcelWritesDisabledError, SafeWorkbookWriter, WorkbookChangedError, WorkbookConflictError, WorkbookLockedError,
)
from vspm.excel.xlsx_package import SheetEditor, XlsxPackage
from tests.conftest import EXTRA_PARTS, FORMULAS


def payload(**kw) -> ExcelRowPayload:
    d = dict(material="ΓΑΛΒΑΝΙΖΕ", thickness="0.60", dimensions="34ΧΡΟΛΟ", pieces="4", weight="496",
             customer="ΓΚΟΥΝΤΑ", category="PG", barcode="PG7030", location="7", roll="E1830460001",
             print_date="2026-10-21")
    d.update(kw)
    return ExcelRowPayload(**d)


def writer(path, tmp_path, **kw) -> SafeWorkbookWriter:
    return SafeWorkbookWriter(path, DEFAULT_SHEET_NAME, backup_dir=tmp_path / "bk", write_enabled=True, **kw)


# ---- formulas ----------------------------------------------------------------------------
def test_formula_row_translation_examples():
    for col, f in FALLBACK_FORMULAS.items():
        at350 = translate_to_row(f.format(r=329).replace("{r}", "329"), 329, 350)
        assert at350 == f.replace("{r}", "350")
    assert "K351" in translate_to_row(FALLBACK_FORMULAS["L"].replace("{r}", "350"), 350, 351)
    assert "F351" in translate_to_row(FALLBACK_FORMULAS["N"].replace("{r}", "350"), 350, 351)


def test_formula_external_links_preserved():
    f = 'IFERROR(INDEX([1]!Table4[ΠΟΙΟΤΗΤΑ],MATCH(K2,[1]!Table4[Νο ΡΟΛΛΟΥ],0)),"-")'
    out = shift_rows(f, 348)
    assert out == f.replace("K2", "K350")
    assert out.count("[1]!Table4") == 2


def test_formula_absolute_strings_and_sheets_untouched():
    f = 'SUM($A$1,A$2,$A3)&"K5"&\'Sheet 1\'!B2&ID!$A:$B'
    assert shift_rows(f, 10) == 'SUM($A$1,A$2,$A13)&"K5"&\'Sheet 1\'!B12&ID!$A:$B'


def test_row_pattern():
    assert row_pattern('IFERROR(VLOOKUP(F7,ID!$A:$B,2,0),"-")', 7) == 'IFERROR(VLOOKUP(F{r},ID!$A:$B,2,0),"-")'


# ---- writer ---------------------------------------------------------------------------------
def test_append_exact_column_mapping(workbook, tmp_path):
    rep = writer(workbook, tmp_path).append([payload()])
    assert rep.results[0].status == "written" and rep.results[0].row_number == 5
    wb = openpyxl.load_workbook(workbook)
    ws = wb[DEFAULT_SHEET_NAME]
    vals = [ws.cell(5, c).value for c in range(1, 17)]
    assert vals[:11] == ["ΓΑΛΒΑΝΙΖΕ", "0.60", "34ΧΡΟΛΟ", 4, 496, "ΓΚΟΥΝΤΑ", "PG", "PG7030", 7, None,
                         "E1830460001"]
    for i, col in enumerate("LMNO"):
        assert vals[11 + i] == FORMULAS[col].format(r=5)
    assert vals[15] == datetime(2026, 10, 21)
    assert ws.cell(5, 16).number_format == "dd\\/mm\\/yyyy"
    # existing rows unchanged
    assert ws.cell(4, 8).value == "VS9002" and ws.cell(2, 10).value == "Α"


def test_consecutive_rows_get_their_own_references(workbook, tmp_path):
    w = writer(workbook, tmp_path)
    w.append([payload(barcode="PG1")])
    w.append([payload(barcode="PG2")])
    ws = openpyxl.load_workbook(workbook)[DEFAULT_SHEET_NAME]
    assert ws.cell(6, 12).value == FORMULAS["L"].format(r=6)
    assert ws.cell(6, 14).value == FORMULAS["N"].format(r=6)
    assert ws.cell(6, 8).value == "PG2"


def test_duplicate_write_is_prevented(workbook, tmp_path):
    w = writer(workbook, tmp_path)
    w.append([payload()])
    rep = w.append([payload()])
    assert rep.results[0].status == "already_present" and rep.results[0].row_number == 5
    ed = SheetEditor(XlsxPackage(workbook.read_bytes()), DEFAULT_SHEET_NAME)
    assert ed.find_rows_with_value("H", "PG7030") == [5]


def test_other_package_parts_preserved_byte_for_byte(workbook, tmp_path):
    before = zipfile.ZipFile(workbook)
    before_parts = {n: before.read(n) for n in before.namelist()}
    writer(workbook, tmp_path).append([payload()])
    after = zipfile.ZipFile(workbook)
    assert after.namelist() == list(before_parts)
    changed = {n for n in before_parts if after.read(n) != before_parts[n]}
    assert changed <= {"xl/workbook.xml", "xl/styles.xml", "xl/worksheets/sheet2.xml"}
    for name, data in EXTRA_PARTS.items():
        assert after.read(name) == data


def test_hidden_preformatted_row_becomes_visible(tmp_path):
    from tests.conftest import build_test_workbook
    wb = build_test_workbook(tmp_path / "w.xlsx", data_rows=3, styled_empty_rows=4)
    ed = SheetEditor(XlsxPackage(wb.read_bytes()), DEFAULT_SHEET_NAME)
    assert ed.rows()[7].get("hidden") == "1"
    w = writer(wb, tmp_path)
    for i in range(4):
        w.append([payload(barcode=f"PG{i}")])
    ws = openpyxl.load_workbook(wb)[DEFAULT_SHEET_NAME]
    assert ws.cell(8, 8).value == "PG3"
    assert not ws.row_dimensions[8].hidden


def test_composite_pieces_and_text_roll_kept_as_text(workbook, tmp_path):
    writer(workbook, tmp_path).append([payload(pieces="2+1", roll="0123", location="ΜΙΚΡΟ")])
    ws = openpyxl.load_workbook(workbook)[DEFAULT_SHEET_NAME]
    assert ws.cell(5, 4).value == "2+1" and ws.cell(5, 11).value == "0123" and ws.cell(5, 9).value == "ΜΙΚΡΟ"


def test_column_j_never_written(workbook, tmp_path):
    writer(workbook, tmp_path).append([payload()])
    ed = SheetEditor(XlsxPackage(workbook.read_bytes()), DEFAULT_SHEET_NAME)
    assert ed.value(5, "J").kind == "empty"


def test_writes_disabled_by_default(workbook, tmp_path):
    w = SafeWorkbookWriter(workbook, DEFAULT_SHEET_NAME, backup_dir=tmp_path / "bk")
    with pytest.raises(ExcelWritesDisabledError):
        w.append([payload()])


def test_excel_owner_file_blocks_write(workbook, tmp_path):
    (workbook.parent / ("~$" + workbook.name[2:])).write_bytes(b"x")
    with pytest.raises(WorkbookLockedError):
        writer(workbook, tmp_path).append([payload()])


def test_onedrive_conflict_copy_blocks_write(workbook, tmp_path):
    (workbook.parent / f"{workbook.stem}-DESKTOP-ABC{workbook.suffix}").write_bytes(b"x")
    with pytest.raises(WorkbookConflictError):
        writer(workbook, tmp_path).append([payload()])


def test_concurrent_change_aborts_write(workbook, tmp_path, monkeypatch):
    w = writer(workbook, tmp_path)
    original = w.read_bytes
    calls = {"n": 0}

    def changing_read():
        calls["n"] += 1
        data = original()
        return data if calls["n"] == 1 else data + b"changed-by-someone-else"

    monkeypatch.setattr(w, "read_bytes", changing_read)
    before = workbook.read_bytes()
    with pytest.raises(WorkbookChangedError):
        w.append([payload()])
    assert workbook.read_bytes() == before  # nothing written


def test_backup_created_and_reconciliation_lookup(workbook, tmp_path):
    w = writer(workbook, tmp_path)
    rep = w.append([payload()])
    assert rep.backup_path.exists() and rep.hash_before != rep.hash_after
    assert w.find_barcodes(["PG7030", "VS9000", "XX1"]) == {"PG7030": [5], "VS9000": [2], "XX1": []}


def test_missing_sheet_reported(workbook, tmp_path):
    w = SafeWorkbookWriter(workbook, "Ετοιμα προϊόντα", backup_dir=tmp_path, write_enabled=True)
    with pytest.raises(Exception, match="not found"):
        w.append([payload()])
