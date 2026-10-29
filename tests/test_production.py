"""Print & Save transactions: statuses, failures, reprints, restart recovery."""

from __future__ import annotations

from datetime import datetime, timezone

import openpyxl
import pytest

from vspm.config import DEFAULT_SHEET_NAME
from vspm.domain.models import ProductionRow, RowStatus, TemplateKey
from vspm.excel.xlsx_package import SheetEditor, XlsxPackage
from vspm.services.context import build_context
from vspm.services.production import AlreadyPrinted, ProductionError, ValidationFailed
from tests.conftest import with_settings


def make_row(ctx, **kw) -> int:
    base = dict(document_id=None, row_index=1, customer_full="ΓΚΟΥΝΤΑ - ΠΑΓΟΥΝΗ", customer_short="ΓΚΟΥΝΤΑ",
                material="ΓΑΛΒΑΝΙΖΕ", thickness="0,60", dimensions="34ΧΡΟΛΟ", weight="496", pieces="4",
                packages="1", source_barcode="B-101", source_roll="E1830460001", location="7",
                template_key=TemplateKey.PG, approved=True, status=RowStatus.approved)
    base.update(kw)
    return ctx.repo.add_row(ProductionRow(**base))


@pytest.fixture
def ctx(settings, clock):
    c = build_context(settings)
    c.production._clock = clock
    return c


def excel_rows(path, barcode):
    return SheetEditor(XlsxPackage(path.read_bytes()), DEFAULT_SHEET_NAME).find_rows_with_value("H", barcode)


def test_print_and_save_happy_path(ctx, settings):
    rid = make_row(ctx)
    out = ctx.production.print_and_save(rid)
    assert out.print_ok and out.barcode == "PG990000"
    assert out.print_status == "print_unconfirmed" and out.excel_status == "synchronized"
    lab = ctx.repo.get_label(out.label_id)
    assert lab["print_date_local"] == "2026-10-21"
    assert excel_rows(settings.excel_path, "PG990000") == [5]
    ctx.production.confirm_print(out.label_id, True)
    assert ctx.repo.get_label(out.label_id)["print_status"] == "printed_confirmed"
    assert ctx.repo.get_row(rid).status == RowStatus.printed_confirmed


def test_reprint_reuses_barcode_and_never_adds_excel_row(ctx, settings):
    rid = make_row(ctx)
    out = ctx.production.print_and_save(rid)
    before = settings.excel_path.read_bytes()
    re = ctx.production.reprint(out.label_id)
    assert re.barcode == out.barcode and re.print_ok
    assert settings.excel_path.read_bytes() == before
    assert len(ctx.printer.printed) == 2 and ctx.printer.printed[0] == ctx.printer.printed[1]
    with pytest.raises(AlreadyPrinted):
        ctx.production.print_and_save(rid)
    kinds = [a["kind"] for a in ctx.repo.print_attempts(out.label_id)]
    assert kinds == ["initial", "reprint"]


def test_unapproved_or_invalid_row_cannot_print(ctx):
    rid = make_row(ctx, approved=False)
    with pytest.raises(ValidationFailed) as e:
        ctx.production.print_and_save(rid)
    assert "approved" in e.value.errors
    rid2 = make_row(ctx, weight="", flags={"pieces": "uncertain"})
    with pytest.raises(ValidationFailed) as e:
        ctx.production.print_and_save(rid2)
    assert {"weight", "pieces"} <= set(e.value.errors)
    assert ctx.repo.label_for_row(rid2) is None  # no barcode consumed


def test_failed_print_keeps_barcode_and_retry_uses_it(ctx, settings):
    rid = make_row(ctx)
    ctx.printer.fail_next = "printer offline"
    out = ctx.production.print_and_save(rid)
    assert not out.print_ok and out.print_status == "error" and out.excel_status == "not_queued"
    assert excel_rows(settings.excel_path, out.barcode) == []
    out2 = ctx.production.print_and_save(rid)  # retry
    assert out2.print_ok and out2.barcode == out.barcode
    assert out2.excel_status == "synchronized"
    assert [a["status"] for a in ctx.repo.print_attempts(out.label_id)] == ["error", "submitted"]


def test_failed_reprint_does_not_downgrade_label(ctx):
    out = ctx.production.print_and_save(make_row(ctx))
    ctx.printer.fail_next = "paper out"
    re = ctx.production.reprint(out.label_id)
    assert not re.print_ok and re.print_status == "print_unconfirmed"


def test_failed_excel_sync_is_recorded_and_recoverable(ctx, settings):
    lock = settings.excel_path.parent / ("~$" + settings.excel_path.name[2:])
    lock.write_bytes(b"x")
    out = ctx.production.print_and_save(make_row(ctx))
    assert out.print_ok and out.excel_status == "conflict"
    assert any("NOT updated" in m for m in out.messages)
    lock.unlink()
    msgs = ctx.production.sync_excel(out.label_id)
    assert "row 5" in msgs[0]
    assert ctx.repo.get_label(out.label_id)["excel_status"] == "synchronized"
    statuses = [a["status"] for a in ctx.repo.excel_attempts(out.label_id)]
    assert statuses == ["conflict", "synchronized"]


def test_excel_write_is_idempotent_on_retry(ctx, settings):
    out = ctx.production.print_and_save(make_row(ctx))
    ctx.repo.update_label(out.label_id, excel_status="error")  # pretend the result was lost
    ctx.production.sync_excel(out.label_id)
    assert excel_rows(settings.excel_path, out.barcode) == [5]
    assert ctx.repo.excel_attempts(out.label_id)[-1]["status"] == "already_present"


def test_registry_unavailable_blocks_printing_without_issuing_numbers(settings, clock):
    s = with_settings(settings, **{"barcode_mode": "registry", "registry_url": "http://127.0.0.1:9",
                                    "registry_api_key": "k"})
    c = build_context(s)
    rid = make_row(c)
    with pytest.raises(ProductionError, match="unreachable"):
        c.production.print_and_save(rid)
    assert c.repo.label_for_row(rid) is None
    assert c.printer.printed == []


def test_barcode_disabled_mode_blocks_printing(settings):
    c = build_context(with_settings(settings, **{"barcode_mode": "disabled"}))
    with pytest.raises(ProductionError, match="disabled"):
        c.production.print_and_save(make_row(c))


def test_restart_recovery_after_crash_before_excel(ctx, settings, clock):
    rid = make_row(ctx)
    out = ctx.production.print_and_save(rid)
    # simulate a crash right after printing, before Excel was queued
    ctx.repo.update_label(out.label_id, excel_status="not_queued", excel_payload_json=None)
    c2 = build_context(settings)
    assert any(out.barcode in n for n in c2.recovery_notes)
    assert c2.repo.get_label(out.label_id)["excel_status"] == "pending"
    c2.production.sync_excel(out.label_id)
    assert excel_rows(settings.excel_path, out.barcode) == [5]


def test_restart_after_reservation_resumes_with_same_number(ctx, settings):
    rid = make_row(ctx)
    row = ctx.repo.get_row(rid)
    spec = ctx.templates.spec(row.template_key)
    res = ctx.production._reserve(row, spec.barcode_prefix, "op", "pc")
    from vspm.services.transforms import build_label_content
    ctx.production._create_label(row, res, build_label_content(row, spec, ctx.transforms, res.barcode), "op", "pc")
    c2 = build_context(settings)  # "restart"
    assert any("reserved" in n for n in c2.recovery_notes)
    out = c2.production.print_and_save(rid)
    assert out.barcode == res.barcode and out.print_ok


def test_staging_mode_and_reconciliation(settings, clock):
    c = build_context(with_settings(settings, **{"excel_mode": "staging"}))
    c.production._clock = clock
    out = c.production.print_and_save(make_row(c))
    assert out.excel_status == "staged"
    assert excel_rows(settings.excel_path, out.barcode) == []
    res = c.production.reconcile()
    assert res["missing"] == [out.barcode]
    c.production.apply_staged()
    assert c.repo.get_label(out.label_id)["excel_status"] == "synchronized"
    assert excel_rows(settings.excel_path, out.barcode) == [5]


def test_staging_export_file(settings, clock, tmp_path):
    c = build_context(with_settings(settings, **{"excel_mode": "staging"}))
    c.production._clock = clock
    c.production.print_and_save(make_row(c))
    n = c.production.export_staging(tmp_path / "out.xlsx")
    ws = openpyxl.load_workbook(tmp_path / "out.xlsx").active
    assert n == 1 and ws.cell(2, 8).value == "PG990000" and ws.cell(2, 7).value == "PG"
    assert ws.cell(2, 16).value == datetime(2026, 10, 21)


def test_print_date_uses_factory_timezone(ctx, settings):
    ctx.production._clock = lambda: datetime(2026, 10, 21, 22, 30, tzinfo=timezone.utc)  # 01:30 in Athens
    out = ctx.production.print_and_save(make_row(ctx))
    assert ctx.repo.get_label(out.label_id)["print_date_local"] == "2026-10-22"
    ws = openpyxl.load_workbook(settings.excel_path)[DEFAULT_SHEET_NAME]
    assert ws.cell(5, 16).value == datetime(2026, 10, 22)


def test_pg0_row_recorded_as_pg_in_excel(ctx, settings):
    out = ctx.production.print_and_save(make_row(ctx, customer_full="ΚΥΡΙΑΚΙΔΗΣ", customer_short="ΚΥΡΙΑΚΙΔΗΣ",
                                                 template_key=TemplateKey.PG0))
    assert out.barcode.startswith("PG0")
    ws = openpyxl.load_workbook(settings.excel_path)[DEFAULT_SHEET_NAME]
    assert ws.cell(5, 7).value == "PG" and ws.cell(5, 8).value == out.barcode and ws.cell(5, 6).value == "ΚΥΡΙΑΚΙΔΗΣ"


def test_printed_row_values_are_locked(ctx):
    rid = make_row(ctx)
    ctx.production.print_and_save(rid)
    row = ctx.repo.get_row(rid)
    row.weight = "999"
    with pytest.raises(ValueError):
        ctx.repo.save_row(row)
    with pytest.raises(ValueError):
        ctx.repo.delete_row(rid)


def test_each_row_prints_individually_with_unique_numbers(ctx, settings):
    ids = [make_row(ctx, row_index=i, weight=str(100 + i)) for i in range(1, 4)]
    out2 = ctx.production.print_and_save(ids[1])
    out1 = ctx.production.print_and_save(ids[0])
    assert out2.barcode != out1.barcode
    assert ctx.repo.label_for_row(ids[2]) is None
    assert excel_rows(settings.excel_path, out2.barcode) == [5]
    assert excel_rows(settings.excel_path, out1.barcode) == [6]


def test_document_flow_with_mock_source(ctx):
    docs = ctx.documents.latest(5)
    assert len(docs) == 5
    rows = ctx.documents.run_ocr(docs[0].doc_id, ctx.ocr_provider())
    assert len(rows) == 3
    # re-reading the same document id/revision is detected as the same document
    again = ctx.documents.latest(5)
    assert again[0].doc_id == docs[0].doc_id and again[0].rows == 3
